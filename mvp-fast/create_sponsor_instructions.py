from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Title
title = doc.add_heading('Google Sheets Integration Setup Instructions', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('This document contains step-by-step instructions for setting up Google Sheets integration for the ClashArena application.')
doc.add_paragraph('')

# Main heading
doc.add_heading('What You Need to Do', 1)
doc.add_paragraph('To complete the Google Sheets integration, please follow these steps:')
doc.add_paragraph('')

# Step 1
doc.add_heading('Step 1: Create Google Cloud Project', 2)
doc.add_paragraph('1. Go to https://console.cloud.google.com/')
p1 = doc.add_paragraph('2. Sign in with your Google account')
doc.add_paragraph('3. Click the project dropdown at the top (next to "Google Cloud")')
doc.add_paragraph('4. Click "New Project"')
doc.add_paragraph('5. Enter project name: ClashArena (or any name you prefer)')
doc.add_paragraph('6. Click "Create"')
doc.add_paragraph('7. Wait a few seconds, then select your new project from the dropdown')
doc.add_paragraph('')

# Step 2
doc.add_heading('Step 2: Enable Google Sheets API', 2)
doc.add_paragraph('1. In the search bar at the top, type: "Google Sheets API"')
doc.add_paragraph('2. Click on "Google Sheets API" from the results')
doc.add_paragraph('3. Click the blue "Enable" button')
doc.add_paragraph('4. Wait for it to enable (usually takes 10-20 seconds)')
doc.add_paragraph('')

# Step 3
doc.add_heading('Step 3: Create Service Account', 2)
doc.add_paragraph('1. In the left sidebar, click "APIs & Services" > "Credentials"')
doc.add_paragraph('2. Click "+ CREATE CREDENTIALS" at the top')
doc.add_paragraph('3. Select "Service account"')
doc.add_paragraph('4. Fill in the form:')
doc.add_paragraph('   • Service account name: clasharena-sheets (or any name)', style='List Bullet')
doc.add_paragraph('   • Service account ID: Will auto-fill (leave as is)', style='List Bullet')
doc.add_paragraph('   • Description: For ClashArena Google Sheets integration (optional)', style='List Bullet')
doc.add_paragraph('5. Click "Create and Continue"')
doc.add_paragraph('6. Skip the optional steps (click "Continue" then "Done")')
doc.add_paragraph('')

# Step 4
doc.add_heading('Step 4: Download JSON Key File', 2)
doc.add_paragraph('1. You should now see your service account in the list')
doc.add_paragraph('2. Click on the service account name (the email address)')
doc.add_paragraph('3. Go to the "Keys" tab')
doc.add_paragraph('4. Click "Add Key" > "Create new key"')
doc.add_paragraph('5. Select "JSON" format')
doc.add_paragraph('6. Click "Create"')
doc.add_paragraph('7. A JSON file will automatically download to your computer')
p4 = doc.add_paragraph('8. Rename the downloaded file to: credentials.json')
p4.runs[0].bold = True
doc.add_paragraph('9. Send this credentials.json file to the developer')
doc.add_paragraph('')

# Step 5
doc.add_heading('Step 5: Create Google Sheet', 2)
doc.add_paragraph('1. Go to https://sheets.google.com/')
doc.add_paragraph('2. Create a new sheet (or open an existing one)')
doc.add_paragraph('3. Name it (e.g., "ClashArena Entries") - this is optional')
doc.add_paragraph('')

# Step 6
doc.add_heading('Step 6: Share Sheet with Service Account', 2)
doc.add_paragraph('1. Click the "Share" button (top right)')
doc.add_paragraph('2. Open the credentials.json file you downloaded')
doc.add_paragraph('3. Find the line that says "client_email" - it looks like:')
code1 = doc.add_paragraph('   "client_email": "clasharena-sheets@your-project.iam.gserviceaccount.com"')
code1.style = 'Intense Quote'
doc.add_paragraph('4. Copy the entire email address (the part after "client_email": " and before ")')
doc.add_paragraph('5. Paste it into the Google Sheet "Share" dialog')
doc.add_paragraph('6. Make sure permissions are set to "Editor"')
doc.add_paragraph('7. Uncheck "Notify people" (you do not need to notify a service account)')
doc.add_paragraph('8. Click "Share"')
doc.add_paragraph('')

# Step 7
doc.add_heading('Step 7: Get Spreadsheet ID', 2)
doc.add_paragraph('1. Look at your Google Sheet URL in the browser')
doc.add_paragraph('2. It will look like:')
code2 = doc.add_paragraph('   https://docs.google.com/spreadsheets/d/1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms/edit')
code2.style = 'Intense Quote'
doc.add_paragraph('3. Copy the long string between /d/ and /edit')
doc.add_paragraph('   In the example above, it is: 1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms')
doc.add_paragraph('4. This is your Spreadsheet ID')
doc.add_paragraph('5. Send this Spreadsheet ID to the developer')
doc.add_paragraph('')

# Summary
doc.add_heading('Summary: What to Send to Developer', 1)
doc.add_paragraph('Please provide the following to complete the integration:')
doc.add_paragraph('1. The credentials.json file (from Step 4)', style='List Bullet')
doc.add_paragraph('2. The Spreadsheet ID (from Step 7)', style='List Bullet')
doc.add_paragraph('')
doc.add_paragraph('Once the developer receives these items, the integration will be completed and entries will automatically be saved to your Google Sheet.')
doc.add_paragraph('')

# Important Notes
doc.add_heading('Important Notes', 1)
doc.add_paragraph('• Keep your credentials.json file secure - do not share it publicly', style='List Bullet')
doc.add_paragraph('• The service account email must have Editor access to the Google Sheet', style='List Bullet')
doc.add_paragraph('• The sheet structure will be automatically created when the first entry is saved', style='List Bullet')
doc.add_paragraph('• Columns will match your existing CSV format: Time, Gamer ID Name, Gamer Email, Whatsapp Number, Region, Kill, Per kill Value, Money, Dollar Rate, Naira Amount, Bank Name, Bank Account Number, Paid', style='List Bullet')

# Save
try:
    doc.save('Google_Sheets_Setup_Instructions.docx')
    print('Word document created successfully: Google_Sheets_Setup_Instructions.docx')
    print('File location: ' + __import__('os').path.abspath('Google_Sheets_Setup_Instructions.docx'))
except Exception as e:
    print(f'Error creating document: {e}')
    import traceback
    traceback.print_exc()

